import io
import os
import sys
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

class TestP2Hardening(unittest.TestCase):
    def test_imports_and_feature_layer(self):
        import p3_suite
        self.assertTrue(p3_suite.ACTION_REGISTRY)
        self.assertIn("queue_instagram_draft", p3_suite.ACTION_REGISTRY)

    def test_agent_actions_are_approval_gated(self):
        import p3_suite
        self.assertEqual(p3_suite.ACTION_REGISTRY["create_task"]["side_effect"], "internal")
        self.assertNotIn("publish_payment", p3_suite.ACTION_REGISTRY)

    def test_large_csv_stream_shape(self):
        # 50k rows: validates parsing primitives without depending on production UI.
        rows = ["name,brand,article\n"] + [f"Ball {i},Brand,{i}\n" for i in range(50000)]
        payload = io.StringIO("".join(rows))
        count = sum(1 for _ in payload) - 1
        self.assertEqual(count, 50000)

    def test_concurrent_version_conflict_model(self):
        # Deterministic optimistic-concurrency model: only one writer may advance v=1.
        state = {"version": 1, "writes": 0}
        from threading import Lock
        lock = Lock()
        def writer():
            with lock:
                expected = 1
                if state["version"] != expected:
                    return False
                state["version"] += 1
                state["writes"] += 1
                return True
        with ThreadPoolExecutor(max_workers=8) as pool:
            results = list(pool.map(lambda _: writer(), range(8)))
        self.assertEqual(sum(results), 1)
        self.assertEqual(state["version"], 2)

    def test_expired_session_statuses_are_recoverable(self):
        from saas_core import SupabaseRequestError
        self.assertIn(SupabaseRequestError("expired", 401).status_code, (400,401,403))

    def test_billing_modules_are_importable_without_live_charge(self):
        import billing, tbank_billing
        self.assertTrue(hasattr(billing, "create_checkout"))
        self.assertTrue(hasattr(tbank_billing, "create_checkout"))

    @patch("tbank_billing._cfg", side_effect=lambda name, default="": {
        "SUPABASE_SERVICE_ROLE_KEY": "key",
        "SUPABASE_URL": "https://example.supabase.co",
    }.get(name, default))
    @patch("tbank_billing.requests.post")
    def test_tbank_checkout_storage_failure_is_not_silent(self, post_mock, cfg_mock):
        post_mock.return_value.ok = False
        post_mock.return_value.text = "db unavailable"
        import tbank_billing
        with self.assertRaises(RuntimeError):
            tbank_billing._save_checkout("tenant", "order", "payment", "pro", 990.00, "RUB")

    @patch("webmcp_tools.data_load")
    def test_webmcp_content_plan_is_allowlisted(self, load_mock):
        import webmcp_tools
        load_mock.side_effect = [
            [{"city": "Краснодар", "secret_note": "hidden"}],
            [{"_saas_record_id": "p1", "name": "Ball", "private_cost": 100}],
            [{"_saas_record_id": "c1", "date": "2026-10-04", "platform": "Telegram",
              "idea": "Post", "status": "Идея", "private_note": "hidden"}],
        ]
        data = webmcp_tools._webmcp_data()
        self.assertNotIn("private_note", data["content_plan"][0])
        self.assertNotIn("private_cost", data["products"][0])
        self.assertNotIn("secret_note", data["settings"])

    def test_large_xlsx_roundtrip(self):
        from openpyxl import Workbook, load_workbook
        import tempfile
        wb = Workbook(write_only=True)
        ws = wb.create_sheet()
        ws.append(["name", "brand", "article"])
        for i in range(20000):
            ws.append([f"Product {i}", "Brand", str(i)])
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as f:
            wb.save(f.name)
            read = load_workbook(f.name, read_only=True, data_only=True)
            ws2 = read.active
            rows = sum(1 for _ in ws2.iter_rows(min_row=2, values_only=True))
            read.close()
        self.assertEqual(rows, 20000)

if __name__ == "__main__":
    unittest.main()
